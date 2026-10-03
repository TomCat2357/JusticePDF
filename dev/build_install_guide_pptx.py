"""Generate JusticePDFの導入方法.pptx (4:3) from scratch.

配布版「JusticePDF-PRO」の導入簡易マニュアル。読者は PC に詳しくない一般職員。
流れは「Google Drive から zip をダウンロード → 展開 → ショートカット作成 → 起動」。

使う画像は ``dev/manual_assets/`` の 16 / 17 / 19（Google Drive 画面と zip の
「すべて展開」操作）。右クリックメニュー（ps1 の「PowerShell で実行」）や
PowerShell 画面など撮影できないものは、偽のスクリーンショットを作らず
図形（角丸ボックス・矢印・メニュー項目名を囲んだ文字）で表現する。

Run (python-pptx 1.0.2 / Pillow / PyMuPDF が入っている scoop の Python):
    C:\\Users\\sa11882\\scoop\\apps\\python311\\current\\python.exe dev\\build_install_guide_pptx.py

出力: <repo>\\JusticePDFの導入方法.pptx （生成物なので .gitignore の *.pptx で管理外）
"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = Path(__file__).resolve().parent / "manual_assets"
OUT_PPTX = ROOT / "JusticePDFの導入方法.pptx"

# ---------- Design tokens ----------

FONT = "Meiryo UI"      # 日本語（latin / ea / cs すべてに設定する）
MONO = "Consolas"       # コマンド表記（日本語部分は FONT にフォールバック）


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


NAVY = _rgb("1F3A5F")
BLUE = _rgb("2F6DB5")
LIGHT_BLUE = _rgb("E8F0FA")
AMBER = _rgb("F2A93B")
LIGHT_AMBER = _rgb("FFF4DC")
RED = _rgb("C8483C")
LIGHT_RED = _rgb("FDECEA")
GREEN = _rgb("2E8B57")
LIGHT_GREEN = _rgb("E6F4EC")
TEXT = _rgb("2B2F36")
GRAY = _rgb("5F6B7A")
LINE = _rgb("C9D3E0")
BG = _rgb("F5F7FA")
WHITE = _rgb("FFFFFF")
DARK = _rgb("1E2530")
PALE = _rgb("DCE6F5")

TOTAL_PAGES = 10
SLIDE_W, SLIDE_H = 10.0, 7.5


# ---------- Low-level helpers ----------

def _apply_font(run, size, bold, color, face) -> None:
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = face  # <a:latin>
    # 日本語は <a:ea> / <a:cs> を明示しないとテーマ既定の東アジアフォントになる
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def _fill_paras(tf, paras, size, color, bold, align, face,
                line_spacing=None, space_after=0) -> None:
    """paras: list of (str | [run, ...] | {"runs": [...], overrides})。
    run は str か (text, {size,bold,color,face})。"""
    first = True
    for p in paras:
        opts = {}
        if isinstance(p, dict):
            opts = p
            p = p["runs"]
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = opts.get("align", align)
        if line_spacing:
            para.line_spacing = line_spacing
        para.space_after = Pt(opts.get("space_after", space_after))
        runs = [p] if isinstance(p, (str, tuple)) else p
        for r in runs:
            text, ro = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = text
            _apply_font(
                run,
                ro.get("size", opts.get("size", size)),
                ro.get("bold", opts.get("bold", bold)),
                ro.get("color", opts.get("color", color)),
                ro.get("face", face),
            )


def _setup_frame(tf, margins, anchor) -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(margins[0])
    tf.margin_top = tf.margin_bottom = Inches(margins[1])
    tf.vertical_anchor = anchor


def add_text(slide, x, y, w, h, paras, size=20, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, face=FONT,
             margins=(0.05, 0.03), line_spacing=None, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _setup_frame(tb.text_frame, margins, anchor)
    _fill_paras(tb.text_frame, paras, size, color, bold, align, face,
                line_spacing, space_after)
    return tb


def add_shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0,
              radius=None, paras=None, size=18, color=TEXT, bold=False,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, face=FONT,
              margins=(0.08, 0.04), dash=False, line_spacing=None,
              space_after=0):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    if radius is not None and len(shp.adjustments) > 0:
        shp.adjustments[0] = radius
    _setup_frame(shp.text_frame, margins, anchor)
    if paras:
        _fill_paras(shp.text_frame, paras, size, color, bold, align, face,
                    line_spacing, space_after)
    return shp


def rrect(slide, x, y, w, h, **kw):
    kw.setdefault("radius", 0.12)
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, **kw)


def circle_num(slide, x, y, d, n, fill=BLUE, size=24, color=WHITE):
    return add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill,
                     paras=[str(n)], size=size, color=color, bold=True,
                     margins=(0, 0))


def arrow_right(slide, x, y, w=0.3, h=0.4, fill=GRAY):
    return add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, fill=fill)


def add_picture(slide, name, x, y, width=None, height=None, crop=None, alt=""):
    """画像を縦横比そのままで配置する。crop は元画像 px の (l, t, r, b)。
    戻り値: (picture, scale[inch/px], crop_left_px, crop_top_px)"""
    path = ASSETS / name
    if not path.exists():
        raise FileNotFoundError(f"画像が見つかりません: {path}")
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, r, b = crop if crop else (0, 0, iw, ih)
    cw, ch = r - l, b - t
    if width is not None:
        scale = width / cw
    else:
        scale = height / ch
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(cw * scale), Inches(ch * scale))
    if crop:
        pic.crop_left = l / iw
        pic.crop_top = t / ih
        pic.crop_right = (iw - r) / iw
        pic.crop_bottom = (ih - b) / ih
    pic.line.color.rgb = LINE
    pic.line.width = Pt(1)
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic, scale, l, t


def highlight(slide, pic_x, pic_y, scale, l, t, px_box, pad_px=6, color=RED,
              line_w=3.0):
    """画像上の px 矩形 (x0, y0, x1, y1) を赤枠で囲む。"""
    x0, y0, x1, y1 = px_box
    return add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        pic_x + (x0 - pad_px - l) * scale, pic_y + (y0 - pad_px - t) * scale,
        (x1 - x0 + 2 * pad_px) * scale, (y1 - y0 + 2 * pad_px) * scale,
        fill=None, line=color, line_w=line_w, radius=0.2)


# ---------- Slide scaffolding ----------

def new_slide(prs, title, page, step=None):
    """Title Only レイアウトにタイトル帯・進捗表示・ページ番号を付けて返す。"""
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG

    band = add_shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, 1.2, fill=NAVY)
    rule = add_shape(s, MSO_SHAPE.RECTANGLE, 0, 1.2, SLIDE_W, 0.06, fill=AMBER)
    # タイトルのプレースホルダより奥（背面）に回す
    tree = s.shapes._spTree
    for shp in (rule, band):
        tree.remove(shp._element)
        tree.insert(2, shp._element)

    ttl = s.shapes.title
    if step:
        tx, tw = 1.5, 6.4
        circle_num(s, 0.45, 0.2, 0.8, step, fill=WHITE, size=36, color=NAVY)
        add_text(s, 1.5, 0.1, 3.0, 0.35, [f"STEP {step}"], size=14,
                 color=AMBER, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        ty, th = 0.42, 0.7
        # 進捗ドット（4 ステップ中どこか）
        for i in range(1, 5):
            add_shape(s, MSO_SHAPE.OVAL, 7.9 + (i - 1) * 0.4, 0.5, 0.24, 0.24,
                      fill=AMBER if i <= step else None,
                      line=AMBER, line_w=1.5)
        add_text(s, 7.7, 0.15, 1.9, 0.32, [f"{step} / 4"], size=14, color=PALE,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    else:
        tx, tw, ty, th = 0.5, 9.0, 0.2, 0.85
    ttl.left, ttl.top = Inches(tx), Inches(ty)
    ttl.width, ttl.height = Inches(tw), Inches(th)
    _setup_frame(ttl.text_frame, (0.05, 0.02), MSO_ANCHOR.MIDDLE)
    _fill_paras(ttl.text_frame, [title], 28, WHITE, True, PP_ALIGN.LEFT, FONT)

    add_text(s, 0.5, 7.15, 5.0, 0.3, ["JusticePDF 導入ガイド"], size=12,
             color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 8.2, 7.15, 1.3, 0.3, [f"{page} / {TOTAL_PAGES}"], size=12,
             color=GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    return s


def step_row(s, x, y, w, h, n, paras, size=20, fill=BLUE):
    """番号の丸 + 本文（縦中央揃え）の 1 行。"""
    d = 0.6
    circle_num(s, x, y + (h - d) / 2, d, n, fill=fill, size=24)
    return add_text(s, x + d + 0.2, y, w - d - 0.2, h, paras, size=size,
                    anchor=MSO_ANCHOR.MIDDLE)


def note_box(s, x, y, w, h, paras, kind="warn", size=18):
    fill, line = {
        "warn": (LIGHT_AMBER, AMBER),
        "ng": (LIGHT_RED, RED),
        "ok": (LIGHT_GREEN, GREEN),
        "info": (LIGHT_BLUE, BLUE),
    }[kind]
    return rrect(s, x, y, w, h, fill=fill, line=line, line_w=1.5,
                 paras=paras, size=size, align=PP_ALIGN.LEFT,
                 margins=(0.2, 0.06), radius=0.1)


def B(text, color=None, **kw):
    """太字 run のショートカット。"""
    o = {"bold": True}
    if color is not None:
        o["color"] = color
    o.update(kw)
    return (text, o)


def C(text, color=NAVY, size=None):
    """等幅（コマンド・パス）run のショートカット。"""
    o = {"face": MONO, "bold": True, "color": color}
    if size:
        o["size"] = size
    return (text, o)


# ---------- Slides ----------

def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])   # Title Only（アウトライン用）
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    add_shape(s, MSO_SHAPE.RECTANGLE, 0, 6.9, SLIDE_W, 0.6, fill=_rgb("182E4D"))

    add_shape(s, MSO_SHAPE.FOLDED_CORNER, 4.15, 0.85, 1.7, 2.1, fill=WHITE,
              paras=[B("PDF", RED)], size=34)
    ttl = s.shapes.title
    ttl.left, ttl.top, ttl.width, ttl.height = (
        Inches(0.5), Inches(3.3), Inches(9.0), Inches(1.0))
    _setup_frame(ttl.text_frame, (0.05, 0.03), MSO_ANCHOR.MIDDLE)
    _fill_paras(ttl.text_frame, ["JusticePDF 導入ガイド"], 44, WHITE, True,
                PP_ALIGN.CENTER, FONT)
    add_shape(s, MSO_SHAPE.RECTANGLE, 4.2, 4.45, 1.6, 0.06, fill=AMBER)
    add_text(s, 0.5, 4.65, 9.0, 0.7, ["ダウンロードから起動まで"], size=28,
             color=PALE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    labels = ["Windows 10 / 11", "インストール不要", "管理者権限は不要"]
    w, gap = 2.6, 0.25
    x0 = (SLIDE_W - (3 * w + 2 * gap)) / 2
    for i, t in enumerate(labels):
        rrect(s, x0 + i * (w + gap), 5.75, w, 0.65, fill=None, line=PALE,
              line_w=1.5, paras=[t], size=18, color=WHITE, radius=0.5)


def slide_flow(prs):
    s = new_slide(prs, "全体の流れ　〜 4 つのステップだけ", 2)
    cards = [
        ("1", "ダウンロード", ["Google Drive", "から保存"]),
        ("2", "展開", ["zip を", "ユーザー", "フォルダに展開"]),
        ("3", "ショートカット", ["デスクトップに", "アイコンを作る"]),
        ("4", "起動", ["アイコンを", "ダブルクリック"]),
    ]
    cw, gap, x0, y0, ch = 1.95, 0.4, 0.5, 1.55, 2.7
    for i, (n, head, body) in enumerate(cards):
        x = x0 + i * (cw + gap)
        rrect(s, x, y0, cw, ch, fill=WHITE, line=LINE, line_w=1.5, radius=0.06)
        circle_num(s, x + (cw - 0.7) / 2, y0 + 0.15, 0.7, n, size=28)
        add_text(s, x, y0 + 0.95, cw, 0.5, [head], size=18, bold=True,
                 color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 margins=(0.02, 0.02))
        add_text(s, x, y0 + 1.5, cw, 1.15, body, size=18, align=PP_ALIGN.CENTER,
                 margins=(0.02, 0.02))
        if i < 3:
            arrow_right(s, x + cw + 0.05, y0 + ch / 2 - 0.2, 0.3, 0.4)

    # 必要なもの
    rrect(s, 0.5, 4.4, 9.0, 2.6, fill=LIGHT_BLUE, line=None, radius=0.05)
    add_text(s, 0.75, 4.5, 4.0, 0.5, ["必要なもの"], size=22, bold=True,
             color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    items = [
        ["Windows 10 または 11 のパソコン"],
        ["空き容量 約 2GB（zip 約 570MB ＋ 展開後 約 1.4GB）"],
        ["管理者権限は不要・インストール作業もなし"],
        ["Python・OCR・辞書もすべて同梱済み"],
    ]
    for i, runs in enumerate(items):
        y = 5.08 + i * 0.47
        add_shape(s, MSO_SHAPE.OVAL, 0.8, y + 0.04, 0.32, 0.32, fill=GREEN,
                  paras=["✓"], size=14, color=WHITE, bold=True, margins=(0, 0))
        add_text(s, 1.3, y, 8.0, 0.4, runs, size=20, anchor=MSO_ANCHOR.MIDDLE)


def slide_step1(prs):
    s = new_slide(prs, "ダウンロードする", 3, step=1)
    px, py = 0.5, 1.45
    pic, sc, l, t = add_picture(
        s, "16_install_zip.png", px, py, width=9.0,
        alt="Google ドライブのファイル一覧。JusticePDFの使い方.docx と "
            "JusticePDF-PRO.zip の 2 行が並んでいる。zip の行を赤枠で囲んでいる")
    # JusticePDF-PRO.zip の行（画像下側。行の下端は画像の下端で切れている）
    highlight(s, px, py, sc, l, t, (32, 131, 1232, 164), pad_px=2)
    # 行の下から指す吹き出し
    add_shape(s, MSO_SHAPE.UP_ARROW, px + 170 * sc - 0.2, py + 168 * sc + 0.05,
              0.4, 0.4, fill=RED)
    rrect(s, 2.0, py + 168 * sc + 0.12, 5.2, 0.8, fill=LIGHT_RED, line=RED,
          line_w=1.5, paras=["この zip ファイルを右クリックして", "「ダウンロード」を選ぶ"],
          size=18, bold=True, color=RED, radius=0.12)

    y = 3.8
    step_row(s, 0.5, y, 9.0, 0.95, 1, [
        ["zip を右クリックして", B("「ダウンロード」", BLUE), "を選ぶ"],
        {"runs": ["（Google Drive の画面で。フォルダごとではなく、zip 単体で）"],
         "size": 18, "color": GRAY},
    ])
    step_row(s, 0.5, y + 1.1, 9.0, 1.3, 2, [
        ["「ウイルス スキャンを実行できません」と出たら"],
        [B("「このままダウンロード」", BLUE), "を押す"],
        {"runs": ["（ファイルが大きいための表示です。心配ありません）"],
         "size": 18, "color": GRAY},
    ])
    step_row(s, 0.5, y + 2.5, 9.0, 0.7, 3, [
        ["ダウンロードが終わるまで待つ（約 570MB）"]])


def slide_step2a(prs):
    s = new_slide(prs, "展開する ①　すべて展開を選ぶ", 4, step=2)
    px, py = 0.6, 1.45
    # 画像全体を使う（右クリックしたメニューと、下端の選択中の zip が両方写っている）
    pic, sc, l, t = add_picture(
        s, "17_install_extract_menu.png", px, py, height=5.0,
        alt="zip ファイルを右クリックしたメニュー。「すべて展開」の項目を赤枠で囲んでいる")
    highlight(s, px, py, sc, l, t, (110, 209, 346, 233), pad_px=1)
    pic_w = 363 * sc
    add_text(s, px + pic_w / 2 - 1.8, py + 5.05, 3.6, 0.35,
             ["（Windows 11 の画面の例）"], size=14, color=GRAY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    x, w = 4.4, 5.1
    step_row(s, x, 1.4, w, 1.2, 1, [
        ["ダウンロードした"],
        ["zip がある場所を開く"],
        {"runs": ["（ふつうは「ダウンロード」フォルダ）"], "size": 16,
         "color": GRAY},
    ])
    step_row(s, x, 2.7, w, 1.0, 2, [
        ["ダウンロードした zip を", B("右クリック", BLUE), "する"]])
    step_row(s, x, 3.9, w, 1.0, 3, [
        ["メニューの", B("「すべて展開」", RED), "を選ぶ"]])
    note_box(s, x, 5.15, w, 1.6, [
        ["次の画面で、展開先の末尾を消し、"],
        [B("ユーザーフォルダ", NAVY), "だけにします"],
        {"runs": ["（仮想環境でも消えずに残る場所です）"], "size": 16,
         "color": GRAY}],
        kind="warn", size=20)


def slide_step2b(prs):
    s = new_slide(prs, "展開する ②　展開先の末尾を消す", 5, step=2)
    px, py = 0.5, 1.45
    pic, sc, l, t = add_picture(
        s, "19_install_extract_button.png", px, py, width=4.4,
        alt="「圧縮 (ZIP 形式) フォルダーの展開」画面。展開先の欄の末尾"
            "「\\Downloads\\JusticePDF-PRO」を赤枠で囲んでいる")
    # 消す部分「\Downloads\JusticePDF-PRO」（選択された文字列の後ろ 25 文字）を囲む
    # 直前の数字「sa11882」の末尾にかからないよう、枠は ¥ の直前から細めに引く
    highlight(s, px, py, sc, l, t, (208, 226, 447, 252), pad_px=0, line_w=2.25)
    # 囲んだ部分を指す吹き出し（ダイアログの空白部分に重ねる）
    # 矢印はチェックボックスの文字（x ≲ 390px）にかからない右端寄りに置く
    add_shape(s, MSO_SHAPE.UP_ARROW, px + 436 * sc - 0.2, py + 262 * sc, 0.4,
              0.4, fill=RED)
    rrect(s, px + 20 * sc, py + 350 * sc + 0.05, 843 * sc, 0.9, fill=RED,
          line=None,
          paras=[{"runs": ["「\\Downloads\\JusticePDF-PRO」"], "size": 16},
                 {"runs": ["を消す"], "size": 20}],
          size=18, bold=True, color=WHITE, radius=0.12)

    x, w = 5.1, 4.4
    step_row(s, x, 1.45, w, 1.45, 1, [
        ["展開先の欄の末尾の"],
        [C("\\Downloads\\JusticePDF-PRO", RED, 16)],
        ["を消す"]])
    step_row(s, x, 2.95, w, 1.0, 2, [
        [B("「展開」", BLUE), "を押す"],
        {"runs": ["（数分かかることがあります）"], "size": 18, "color": GRAY},
    ])
    step_row(s, x, 4.0, w, 1.15, 3, [
        [C("JusticePDF-PRO", NAVY, 18), " フォルダが"],
        ["できたら完了"]])

    y = 5.3
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, y, 0.7, 0.4, fill=GREEN,
              paras=["OK"], size=18, bold=True, color=WHITE, radius=0.3,
              margins=(0, 0))
    add_text(s, 1.35, y - 0.03, 8.15, 0.46,
             [[C("C:\\Users\\（s付き職員番号）\\JusticePDF-PRO", TEXT, 18),
               "（この形が正解）"]],
             size=18, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.46
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, y, 0.7, 0.4, fill=RED,
              paras=["NG"], size=18, bold=True, color=WHITE, radius=0.3,
              margins=(0, 0))
    add_text(s, 1.35, y - 0.03, 8.15, 0.46,
             [[C("JusticePDF-PRO\\JusticePDF-PRO", TEXT, 18),
               "（二重）→ 展開し直す"]],
             size=18, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.46
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, y, 0.7, 0.4, fill=RED,
              paras=["NG"], size=18, bold=True, color=WHITE, radius=0.3,
              margins=(0, 0))
    add_text(s, 1.35, y - 0.03, 8.15, 0.46,
             ["Downloads・デスクトップ・ドキュメントの中"],
             size=18, color=RED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.5
    add_text(s, 0.5, y, 9.0, 0.35,
             ["（s付き職員番号）は自分の職員番号です（例: s12345）。"
              "入力欄に入っているので残します"],
             size=14, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)


def slide_step3(prs):
    s = new_slide(prs, "ショートカットを作る", 6, step=3)
    step_row(s, 0.5, 1.4, 9.0, 1.5, 1, [
        [C("JusticePDF-PRO", NAVY, 20), " フォルダを開く"],
        {"runs": ["展開が終わると自動で開きます。"], "size": 16, "color": GRAY},
        {"runs": ["開き直すときは、エクスプローラーのアドレスバーに"],
         "size": 16, "color": GRAY},
        {"runs": [C("%USERPROFILE%\\JusticePDF-PRO", BLUE, 18),
                  " と入力して Enter"], "size": 16, "color": GRAY},
    ])

    step_row(s, 0.5, 3.0, 9.0, 1.0, 2, [
        [C("create_desktop_shortcuts.ps1", NAVY, 18), " のファイルを"],
        ["右クリックして、", B("「PowerShell で実行」", BLUE), "を選ぶ"]])
    # メニュー選択の流れ（実際のメニューではなく、項目名を囲んだ図）
    y, h = 4.15, 0.6
    rrect(s, 1.3, y, 1.5, h, fill=WHITE, line=GRAY, line_w=1.5,
          paras=["右クリック"], size=18, bold=True, color=TEXT, radius=0.3)
    arrow_right(s, 2.85, y + 0.1, 0.3, 0.4)
    rrect(s, 3.2, y, 3.3, h, fill=WHITE, line=GRAY, line_w=1.5, dash=True,
          paras=["その他のオプションを確認"], size=18, color=TEXT, radius=0.3)
    add_text(s, 3.2, y + h + 0.02, 3.3, 0.32, ["（Windows 11 のとき）"],
             size=14, color=GRAY, align=PP_ALIGN.CENTER)
    arrow_right(s, 6.55, y + 0.1, 0.3, 0.4)
    rrect(s, 6.9, y, 2.6, h, fill=BLUE, line=None,
          paras=["PowerShell で実行"], size=18, bold=True, color=WHITE,
          radius=0.3)

    step_row(s, 0.5, 5.15, 9.0, 1.0, 3, [
        ["青（または黒）い画面が一瞬出て、すぐ消えます。"],
        ["デスクトップに", B("「JusticePDF」", BLUE),
         "のアイコンができれば完成です。"],
    ])
    note_box(s, 0.5, 6.3, 9.0, 0.65, [
        ["うまくいかないときは、次のページ以降の「困ったとき」を見てください"]],
        kind="warn", size=18)


def slide_step4(prs):
    s = new_slide(prs, "起動する", 7, step=4)
    # 1. ダブルクリック
    circle_num(s, 0.5, 1.6, 0.6, 1, size=24)
    rrect(s, 1.3, 1.55, 2.1, 0.7, fill=BLUE, line=None, paras=["JusticePDF"],
          size=20, bold=True, color=WHITE, radius=0.2)
    add_text(s, 3.5, 1.5, 6.0, 0.8,
             [[B("をダブルクリック", NAVY, size=22)]],
             anchor=MSO_ANCHOR.MIDDLE, size=22)
    add_text(s, 1.3, 2.3, 8.2, 0.5,
             ["（デスクトップのアイコン。Ctrl + Alt + J のキーでも起動できます）"],
             size=18, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)

    # 2. 黒い画面は閉じない（図）
    circle_num(s, 0.5, 3.45, 0.6, 2, size=24)
    # アプリの窓
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.3, 3.2, 2.3, 1.9, fill=WHITE,
              line=GRAY, line_w=1.5, radius=0.06)
    add_shape(s, MSO_SHAPE.RECTANGLE, 1.35, 3.25, 2.2, 0.3, fill=NAVY,
              paras=["JusticePDF"], size=12, color=WHITE, bold=True,
              align=PP_ALIGN.LEFT, margins=(0.08, 0))
    add_text(s, 1.35, 3.7, 1.5, 0.9, ["アプリ", "の画面"], size=18, color=NAVY,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             margins=(0.02, 0.02))
    # 黒い窓（手前）
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.9, 3.9, 2.3, 1.55, fill=DARK,
              line=RED, line_w=2.5, radius=0.06)
    add_shape(s, MSO_SHAPE.RECTANGLE, 2.95, 3.95, 2.2, 0.3, fill=_rgb("49515E"))
    add_text(s, 4.62, 3.95, 0.5, 0.3, ["×"], size=16, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0))
    add_shape(s, MSO_SHAPE.NO_SYMBOL, 4.57, 3.8, 0.6, 0.6, fill=RED)
    add_text(s, 2.9, 4.4, 2.3, 0.9, ["黒い", "コマンド画面"], size=18,
             color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    # 説明
    add_text(s, 5.5, 3.3, 4.0, 2.3, [
        {"runs": ["アプリと一緒に黒い画面も出ます。"], "space_after": 6},
        {"runs": [B("閉じないでください。", RED, size=22)], "space_after": 6},
        ["アプリを閉じると、自動で消えます。"],
    ], size=20, anchor=MSO_ANCHOR.MIDDLE)

    # 3. 初回は待つ
    step_row(s, 0.5, 5.75, 9.0, 1.2, 3, [
        ["初回は起動に少し時間がかかります。"],
        ["画面が出るまで、そのままお待ちください。"],
    ])


def slide_update(prs):
    s = new_slide(prs, "新しい版が配られたら（更新）", 8)
    rows = [
        ["JusticePDF を", B("閉じる", BLUE), "（黒い画面も一緒に消えます）"],
        ["ユーザーフォルダの ", C("JusticePDF-PRO", NAVY, 20), " フォルダを",
         B("削除", BLUE), "する"],
        ["新しい zip を、", B("STEP 1 〜 2", BLUE), " と同じ手順で、",
         "ユーザーフォルダに展開する"],
        [B("STEP 3", BLUE), " と同じ手順で、ショートカットを作り直す"],
    ]
    y = 1.5
    for i, runs in enumerate(rows, start=1):
        step_row(s, 0.5, y, 9.0, 0.95, i, [runs])
        y += 1.1
    note_box(s, 0.5, 6.0, 9.0, 0.95, [
        [B("※仮想デスクトップでは、既定の保存場所（ドキュメント\\PDFs）の"
           "ファイルが消えてしまうことがあります。", RED, size=18)],
        ["大切な PDF は、Google ドライブに保存してください。"],
    ], kind="warn", size=18)


def slide_trouble_shortcut(prs):
    s = new_slide(prs, "困ったとき ①　ショートカットが作れない", 9)
    # 1
    step_row(s, 0.5, 1.45, 9.0, 1.0, 1, [
        [C("JusticePDF-PRO", NAVY, 18), " フォルダを開く",
         ("（開き方は STEP 3 と同じ）", {"size": 16, "color": GRAY})],
        ["上のアドレスバーに ", C("powershell", BLUE, 20),
         " と入力して Enter を押す"]])
    # アドレスバーの図
    rrect(s, 1.3, 2.5, 4.2, 0.5, fill=WHITE, line=GRAY, line_w=1.5,
          paras=[[C("powershell", NAVY, 18)]], size=18, align=PP_ALIGN.LEFT,
          radius=0.2, margins=(0.2, 0.02))
    arrow_right(s, 5.65, 2.55, 0.3, 0.4)
    rrect(s, 6.05, 2.5, 1.2, 0.5, fill=_rgb("E4E8EE"), line=GRAY, line_w=1.5,
          paras=["Enter"], size=18, bold=True, color=TEXT, radius=0.2)
    add_text(s, 1.3, 3.0, 6.0, 0.32,
             ["（アドレスバー：フォルダの場所が表示されている欄）"], size=14,
             color=GRAY, anchor=MSO_ANCHOR.MIDDLE)

    # 2
    step_row(s, 0.5, 3.55, 9.0, 0.8, 2, [
        ["青い画面が開いたら、次の 1 行をコピーして貼り付け、",
         "Enter を押す"]])
    rrect(s, 1.3, 4.45, 8.2, 1.05, fill=DARK, line=None, radius=0.08,
          paras=[[("powershell -NoProfile -ExecutionPolicy Bypass "
                   "-File .\\create_desktop_shortcuts.ps1",
                   {"face": MONO, "bold": True})]],
          size=16, color=WHITE, align=PP_ALIGN.LEFT, margins=(0.25, 0.08))
    # 3
    step_row(s, 0.5, 5.8, 9.0, 0.9, 3, [
        ["デスクトップに", B("「JusticePDF」", BLUE), "ができれば完成です"]])


def slide_trouble_other(prs):
    s = new_slide(prs, "困ったとき ②　こんなときは", 10)
    rows = [
        ("「実行しますか？」などの\n警告が出た", 1.45,
         [[B("「実行」", BLUE), "を押して進めます。"],
          {"runs": ["出にくくする方法：展開の前に、zip を右クリック"],
           "size": 18, "color": GRAY},
          {"runs": ["→「プロパティ」→「許可する」にチェック"],
           "size": 18, "color": GRAY}]),
        ("フォルダを移動したら\n起動しなくなった", 1.0,
         [["ショートカットを", B("作り直します", BLUE)],
          {"runs": ["（STEP 3 と同じ手順）"], "size": 18, "color": GRAY}]),
        ("起動しない", 1.75,
         [[C("run_justice_gui.cmd", NAVY, 18), " を直接"],
          [B("ダブルクリック", BLUE), "して試す"],
          {"runs": ["（JusticePDF-PRO フォルダの中にあります）"], "size": 18,
           "color": GRAY},
          ["それでもだめなら、管理者へ連絡してください"]]),
    ]
    y = 1.45
    for i, (symptom, h, body) in enumerate(rows, start=1):
        rrect(s, 0.5, y, 9.0, h, fill=WHITE, line=LINE, line_w=1.5, radius=0.06)
        circle_num(s, 0.65, y + (h - 0.55) / 2, 0.55, i, fill=NAVY, size=22)
        add_text(s, 1.35, y, 2.7, h,
                 [{"runs": [ln]} for ln in symptom.split("\n")], size=18,
                 bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, 4.1, y, 5.3, h, body, size=18, anchor=MSO_ANCHOR.MIDDLE,
                 space_after=2)
        y += h + 0.1
    note_box(s, 0.5, y + 0.05, 9.0, 0.85, [
        [B("使い方は：", NAVY, size=20),
         " 同梱の ", B("「JusticePDFの使い方.docx」", BLUE, size=20),
         " を見てください"]],
        kind="info", size=20)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)    # 4:3 を明示
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "JusticePDF 導入ガイド"
    prs.core_properties.subject = "ダウンロードから起動まで"
    # python-pptx 既定テンプレートの作者名が残らないようにする
    prs.core_properties.author = "JusticePDF"
    prs.core_properties.last_modified_by = "JusticePDF"

    slide_cover(prs)
    slide_flow(prs)
    slide_step1(prs)
    slide_step2a(prs)
    slide_step2b(prs)
    slide_step3(prs)
    slide_step4(prs)
    slide_update(prs)
    slide_trouble_shortcut(prs)
    slide_trouble_other(prs)
    assert len(prs.slides) == TOTAL_PAGES

    prs.save(str(OUT_PPTX))
    return OUT_PPTX


if __name__ == "__main__":
    print(f"Wrote {build()}")
